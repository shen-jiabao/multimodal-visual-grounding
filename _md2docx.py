#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""技术文档.md → Word（提交材料用）。

与 `_md2pdf.py` 共用同一套块切分口径：标题 / 表格 / 代码块 / 引用 / 列表 / 段落，
并保留行内 **加粗**（转为真正加粗的 run）与 `代码`（转为黑体行内样式），
使 Word 版与 PDF 版内容一致。

用法：python _md2docx.py 技术文档.md 输出.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "等线"
CODE_FONT = "黑体"
NAVY = RGBColor(0x1C, 0x34, 0x66)
GREY = RGBColor(0x66, 0x66, 0x66)

_ITEM_RE = re.compile(r"^([-·*]|\d+\.)\s")
_TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]*`)")
_CJK_RE = re.compile(r"[\u3000-\u303f\u4e00-\u9fff\uff00-\uffef]")


def unescape(s: str) -> str:
    """还原 Markdown 转义字符。"""
    for a, b in (("\\[", "["), ("\\]", "]"), ("\\~", "~"), ("\\|", "|"),
                 ("\\*", "*"), ("\\_", "_"), ("\\@", "@")):
        s = s.replace(a, b)
    return s


def set_font(run, name: str, size: float | None = None,
             bold: bool | None = None, color=None) -> None:
    """设置 run 的字体（含中文字形）与可选字号/加粗/颜色。"""
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_runs(par, text: str, size: float = 10.5,
             bold: bool = False, color=None) -> None:
    """把 **加粗** 与 `代码` 行内标记写成对应的 run。"""
    for seg in _TOKEN_RE.split(unescape(text)):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            set_font(par.add_run(seg[2:-2]), BODY_FONT, size, True, color)
        elif seg.startswith("`") and seg.endswith("`") and len(seg) > 2:
            set_font(par.add_run(seg[1:-1]), CODE_FONT, size - 0.5, bold,
                     color)
        else:
            set_font(par.add_run(seg), BODY_FONT, size, bold, color)


def join_soft(parts: list[str]) -> str:
    """合并软换行（中文间不插空格，拉丁文间补一个空格）。"""
    out = parts[0].strip() if parts else ""
    for nxt in parts[1:]:
        nxt = nxt.strip()
        if not nxt:
            continue
        sep = "" if (_CJK_RE.search(out[-1:]) or _CJK_RE.search(nxt[:1])) \
            else " "
        out = out + sep + nxt
    return out


def read_blocks(lines: list[str]) -> list[tuple[str, object]]:
    """把 Markdown 行切成逻辑块（与 _md2pdf.py 同口径）。"""
    blocks: list[tuple[str, object]] = []
    i, n = 0, len(lines)
    stop = ("#", "|", ">", "```")

    while i < n:
        s = lines[i].strip()
        if s.startswith("```"):
            info = s[3:].strip().lower()
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            blocks.append(("flow" if info == "flow" else "code", buf))
        elif s.startswith("|"):
            buf = []
            while i < n and lines[i].strip().startswith("|"):
                buf.append(lines[i])
                i += 1
            blocks.append(("table", buf))
        elif not s:
            i += 1
        elif s.startswith("#"):
            blocks.append(("head", s))
            i += 1
        elif s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">"))
                i += 1
            blocks.append(("quote", join_soft(buf)))
        else:
            buf = [s]
            i += 1
            while (i < n and lines[i].strip()
                   and not lines[i].strip().startswith(stop)
                   and not _ITEM_RE.match(lines[i].strip())):
                buf.append(lines[i])
                i += 1
            kind = "item" if _ITEM_RE.match(buf[0]) else "para"
            blocks.append((kind, join_soft(buf)))
    return blocks


def add_table(doc, rows: list[list[str]]) -> None:
    """插入 GFM 表格：网格样式、表头加粗、内容居中。"""
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(ncol):
            cell = table.cell(ri, ci)
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            raw = row[ci] if ci < len(row) else ""
            add_runs(par, raw.strip(),
                     size=9, bold=(ri == 0),
                     color=NAVY if ri == 0 else None)
    doc.add_paragraph()


def add_flow(doc, buf: list[str]) -> None:
    """流程图框图：每步一个带边框的单格表，步间插入居中箭头。"""
    steps = [ln.strip() for ln in buf
             if ln.strip() and ln.strip() not in ("▼", "↓", "|", "v")]
    if not steps:
        return
    for i, txt in enumerate(steps):
        table = doc.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        fill = "DEE7F3" if i in (0, len(steps) - 1) else "EFF3F9"
        cell._tc.get_or_add_tcPr().append(
            parse_xml("<w:shd {0} w:val='clear' w:color='auto' "
                      "w:fill='{1}'/>".format(nsdecls("w"), fill)))
        cell.width = Cm(10)
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(par, txt, 10, True, NAVY)
        if i < len(steps) - 1:
            arrow = doc.add_paragraph()
            arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
            arrow.paragraph_format.space_before = Pt(0)
            arrow.paragraph_format.space_after = Pt(0)
            set_font(arrow.add_run("▼"), CODE_FONT, 9, True, NAVY)
    doc.add_paragraph()


def main() -> int:
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2])
    lines = src.read_text(encoding="utf-8").splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10.5)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)

    first_h1 = True
    for kind, payload in read_blocks(lines):
        if kind == "code":
            for ln in payload:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(0)
                set_font(par.add_run(ln if ln.strip() else " "), CODE_FONT, 8.5)
        elif kind == "flow":
            add_flow(doc, payload)
        elif kind == "table":
            rows = []
            for raw in payload:
                cells = re.split(r"(?<!\\)\|", raw.strip().strip("|"))
                if not all(re.fullmatch(r":?-{3,}:?", c.strip())
                           for c in cells):
                    rows.append(cells)
            add_table(doc, rows)
        elif kind == "head":
            level = len(payload) - len(payload.lstrip("#"))
            par = doc.add_paragraph()
            if level == 1 and first_h1:
                first_h1 = False
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(par, payload.lstrip("#").strip(), 18, True, NAVY)
            else:
                size = {2: 14, 3: 12, 4: 11}.get(level, 12)
                add_runs(par, payload.lstrip("#").strip(), size, True, NAVY)
        elif kind == "quote":
            par = doc.add_paragraph()
            add_runs(par, payload, 9.5, False, GREY)
        elif kind == "item":
            mark = re.match(r"^(\d+\.)", payload)
            body = _ITEM_RE.sub("", payload, count=1)
            par = doc.add_paragraph(style="List Bullet" if not mark
                                    else "List Number")
            add_runs(par, body)
        else:
            par = doc.add_paragraph()
            add_runs(par, payload)

    doc.save(str(dst))
    print(f"已写 {dst}（{dst.stat().st_size / 1024:.0f} KB）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
